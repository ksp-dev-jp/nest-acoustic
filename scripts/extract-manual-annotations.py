import json
import sys

import openpyxl


source_path, output_path = sys.argv[1:3]
sheet = openpyxl.load_workbook(source_path, data_only=True).active

recording_date = None
location = None
start_hour = 0
annotations = []

for row in range(4, sheet.max_row + 1):
    recording_date = sheet.cell(row, 2).value or recording_date
    location = sheet.cell(row, 3).value or location
    species = sheet.cell(row, 11).value
    if not species:
        continue

    value = sheet.cell(row, 4).value
    if value is not None:
        start_hour = float(value)
    start_minute = float(sheet.cell(row, 5).value or 0)
    start_second = float(sheet.cell(row, 6).value or 0)
    start = start_hour * 3600 + start_minute * 60 + start_second

    end_values = [sheet.cell(row, column).value for column in (8, 9, 10)]
    if any(value is not None for value in end_values):
        end_hour = float(end_values[0] if end_values[0] is not None else start_hour)
        end_minute = float(end_values[1] if end_values[1] is not None else start_minute)
        end_second = float(end_values[2] if end_values[2] is not None else start_second + 3)
        end = end_hour * 3600 + end_minute * 60 + end_second
    else:
        end = start + 3

    annotations.append({
        "recording_date": str(recording_date),
        "location": location,
        "start_seconds": start,
        "end_seconds": max(start + 0.1, end),
        "species": str(species).strip(),
        "call_type": str(sheet.cell(row, 12).value or "").strip(),
        "note": str(sheet.cell(row, 13).value or "").strip(),
        "source_row": row,
    })

with open(output_path, "w", encoding="utf-8") as output:
    json.dump(annotations, output, ensure_ascii=False, indent=2)

print(json.dumps({"annotations": len(annotations), "output": output_path}, ensure_ascii=False))
